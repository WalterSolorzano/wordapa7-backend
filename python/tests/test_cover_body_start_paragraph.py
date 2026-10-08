"""Regresión portada: body_start en espacio de PÁRRAFO y fecha no-textbox.

Evidencia origen: archivo real de estudio del trabajo. `detect_cover_ooxml`
guardaba el índice del nodo XML (297) en `body_start_paragraph_idx`, que el
generator consume como índice de `doc.paragraphs`: borraba Resumen,
Introducción y medio cuerpo. Además la fecha vivía en un párrafo normal y se
perdía cuando ya se habían detectado autores por textbox.
"""

import io
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402

from parsing.docx_parser import parse_docx_bytes  # noqa: E402
from parsing.ooxml_cover_detector import detect_cover_ooxml  # noqa: E402


def _docx_with_textbox() -> bytes:
    """Portada con un textbox de autor + fecha como párrafo normal."""
    d = Document()
    d.add_paragraph("UNIVERSIDAD NACIONAL")
    d.add_paragraph("Facultad de Ingenieria")
    host = d.add_paragraph()
    NS = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    )
    ac = parse_xml(
        f'<mc:AlternateContent {NS}>'
        "<mc:Choice Requires=\"wps\"><w:drawing><wp:anchor distT=\"0\" distB=\"0\" distL=\"0\" distR=\"0\" "
        "simplePos=\"0\" relativeHeight=\"1\" behindDoc=\"0\" locked=\"0\" layoutInCell=\"1\" allowOverlap=\"1\">"
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="3600000" cy="900000"/><wp:wrapNone/>'
        '<wp:docPr id="7" name="TxBoxAutor"/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>Br. Autor Uno</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Carnet: 2022-0001</w:t></w:r></w:p></w:txbxContent>"
        "</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:shape xmlns:v=\"urn:schemas-microsoft-com:vml\" style=\"width:283pt;height:70pt\">"
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>Br. Autor Uno</w:t></w:r></w:p>'
        "<w:p><w:r><w:t>Carnet: 2022-0001</w:t></w:r></w:p></w:txbxContent></v:textbox>"
        "</v:shape></w:pict></mc:Fallback></mc:AlternateContent>"
    )
    host._p.append(ac)
    d.add_paragraph("25 de junio del a\u00f1o 2025")
    # índice de párrafo del cuerpo = 4
    d.add_heading("Resumen", level=1)
    d.add_paragraph("Cuerpo del trabajo con texto suficiente para no ser portada.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _parse(raw: bytes, tmp_path, name="t.docx"):
    p = tmp_path / name
    p.write_bytes(raw)
    storage = tmp_path / "storage"
    storage.mkdir(exist_ok=True)
    doc = parse_docx_bytes(raw, name, "sess", storage, skip_page_layout=True)
    return doc, p


def test_body_start_apunta_a_parrafo_real_no_a_nodo_xml(tmp_path):
    raw = _docx_with_textbox()
    doc, p = _parse(raw, tmp_path)
    detect_cover_ooxml(doc, str(p))
    idx = doc.portada["body_start_paragraph_idx"]
    # El cuerpo ("Resumen") está en el párrafo 4 de doc.paragraphs.
    assert idx == 4, f"body_start={idx} (esperado 4, en espacio de párrafo)"
    # No debe borrar el cuerpo: el párrafo apuntado es el heading del cuerpo.
    assert Document(str(p)).paragraphs[idx].text.strip() == "Resumen"


def test_fecha_en_parrafo_normal_se_infiere_con_textbox(tmp_path):
    raw = _docx_with_textbox()
    doc, _ = _parse(raw, tmp_path)
    fields = doc.portada.get("fields", {})
    assert fields.get("date") == "25 de junio del a\u00f1o 2025", fields
    # El carnet del textbox también debe llegar al autor.
    assert "2022-0001" in (fields.get("author") or ""), fields.get("author")
