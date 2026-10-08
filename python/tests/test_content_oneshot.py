"""Tests de la ruta one-shot archivos -> .docx APA 7 (content.oneshot)."""

import docx
import pytest

openpyxl = pytest.importorskip("openpyxl")

from content.oneshot import build_from_files  # noqa: E402


def _make_docx(path, text="Contenido de prueba del documento."):
    d = docx.Document()
    d.add_heading("Introducción", level=1)
    d.add_paragraph(text)
    d.save(str(path))
    return path


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["Col A", "Col B"])
    ws.append(["1", "x"])
    ws.append(["2", "y"])
    wb.save(str(path))
    return path


def _docx_text(path):
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_xlsx_only_con_portada_uni(tmp_path):
    xlsx = _make_xlsx(tmp_path / "datos.xlsx")
    result = build_from_files(
        xlsx_path=xlsx,
        cover_mode="generate_uni_cover",
        title="Mi Tesis",
        storage_dir=tmp_path,
        try_com=False,
    )
    out = result["path"]
    assert out and out.endswith(".docx")
    import os

    assert os.path.exists(out)
    text = _docx_text(out)
    # Portada UNI (F3): lugar por defecto.
    assert "Managua" in text
    # La tabla del Excel llegó al documento.
    assert "Col A" in text and "1" in text


def test_docx_mas_xlsx(tmp_path):
    src = _make_docx(tmp_path / "base.docx")
    xlsx = _make_xlsx(tmp_path / "datos.xlsx")
    result = build_from_files(
        docx_path=src,
        xlsx_path=xlsx,
        storage_dir=tmp_path,
        try_com=False,
    )
    import os

    assert os.path.exists(result["path"])
    text = _docx_text(result["path"])
    assert "Introducción" in text
    assert "Col A" in text


def test_docx_inexistente_lanza(tmp_path):
    with pytest.raises(FileNotFoundError):
        build_from_files(docx_path=tmp_path / "no_existe.docx", storage_dir=tmp_path, try_com=False)
