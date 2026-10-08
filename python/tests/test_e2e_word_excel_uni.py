"""E2E (F5): Word + Excel -> .docx con portada UNI y formato APA 7.

Verifica el camino real del MCP one-shot: `build_from_files` con un .docx y un
.xlsx existentes debe producir un documento con portada UNI, cuerpo con
interlineado doble y sangria de primera linea, y la tabla del Excel.
"""

import docx
import pytest
from docx.shared import Inches

openpyxl = pytest.importorskip("openpyxl")

from content.oneshot import build_from_files  # noqa: E402


def _make_docx(path):
    d = docx.Document()
    d.add_heading("Introducción", level=1)
    d.add_paragraph(
        "Este es un parrafo de cuerpo que debe llevar sangria de primera linea "
        "e interlineado doble segun APA 7."
    )
    d.add_paragraph("Segundo parrafo del cuerpo para verificar el formato APA 7.")
    d.save(str(path))
    return path


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["Col A", "Col B"])
    ws.append(["1", "x"])
    ws.append(["2", "y"])
    wb.save(str(path))
    return path


def test_e2e_word_excel_portada_uni_formato(tmp_path):
    src = _make_docx(tmp_path / "base.docx")
    xlsx = _make_xlsx(tmp_path / "datos.xlsx")
    out = tmp_path / "final.docx"

    result = build_from_files(
        docx_path=src,
        xlsx_path=xlsx,
        cover_mode="generate_uni_cover",
        title="Tesis E2E",
        out_path=out,
        storage_dir=tmp_path,
        try_com=False,
    )

    assert result["path"], "build_from_files no devolvio ruta"
    d = docx.Document(str(result["path"]))
    joined = "\n".join(p.text for p in d.paragraphs)

    # Portada UNI generada sin documento original.
    assert "Managua" in joined
    assert "Tesis E2E" in joined

    # La portada NO debe llevar sangria de primera linea.
    for p in d.paragraphs:
        if "Managua" in p.text or "Tesis E2E" in p.text:
            fli = p.paragraph_format.first_line_indent
            assert fli is None or int(fli) == 0, f"portada con sangria: {fli}"

    # Cuerpo: interlineado doble (via estilo Normal) y sangria de primera linea.
    assert d.styles["Normal"].paragraph_format.line_spacing == 2.0
    body = [
        p
        for p in d.paragraphs
        if "parrafo de cuerpo" in p.text or "Segundo parrafo" in p.text
    ]
    assert body, "no se encontraron parrafos de cuerpo"
    for p in body:
        fli = p.paragraph_format.first_line_indent
        assert fli is not None and abs(int(fli) - int(Inches(0.5))) < 1000, (
            f"sangria de cuerpo inesperada: {fli}"
        )

    # La tabla del Excel llego al documento.
    tabla_excel = [
        t
        for t in d.tables
        if any("Col A" in c.text for c in t.rows[0].cells)
    ]
    assert tabla_excel, "falta la tabla proveniente del Excel"
    filas = [[c.text for c in row.cells] for row in tabla_excel[0].rows]
    assert filas[0] == ["Col A", "Col B"]
    assert ["1", "x"] in filas and ["2", "y"] in filas

    # Tamano de pagina carta (8.5 x 11 in).
    sec = d.sections[0]
    assert abs(int(sec.page_width) - int(Inches(8.5))) < 1000
    assert abs(int(sec.page_height) - int(Inches(11))) < 1000
