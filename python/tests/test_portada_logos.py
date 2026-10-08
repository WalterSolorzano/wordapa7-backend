"""
WordAPA7 - los logos son varios, y elegir UNAN no da el logo de la UNI.

El defecto: `_resolve_logo_path()` NO TENIA ARGUMENTO y siempre terminaba en
`logo_uni.png`. Con el `if LOGO_PATH.exists()` que la acompanaba, que solo salta
si el archivo no esta en disco, elegir UNAN producia un `.docx` con el logo de la
UNI sin decir nada. Es la forma peor de fallar: no hay error, hay un documento
equivocado.

Y el `Cm(5.2)` del logo era un ancho absoluto elegido a ojo. Con Carta y A4
elegibles, un milimetro absoluto se ve distinto en cada hoja: el mismo diseno
salia de dos tamanios. La fraccion del ancho util es lo que lo arregla, y es lo
que mide la tercera prueba.

`_usa_asset` y `_ancho_de_imagen_cm` son helpers de ESTE archivo a proposito: leen
el `.docx` generado, no un diccionario en memoria. Si el generador no escribiera
el asset o no escribiera el tamano, esto no lo veria.
"""

import hashlib
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import docx  # noqa: E402
import pytest  # noqa: E402

from generation.generator import generate_apa7_docx  # noqa: E402
from models import APARuleSet, DocumentModel, ElementModel, ElementType, LogoPortada, PortadaData  # noqa: E402
from modules import portada_uni  # noqa: E402

EMU_POR_CM = 360000


# ── Los helpers ──────────────────────────────────────────────────────────────

def _modelo(session: str) -> DocumentModel:
    return DocumentModel(
        session_id=session,
        file_name="logos.docx",
        elements=[
            ElementModel(id="e1", type=ElementType.PARAGRAPH, text="Cuerpo."),
        ],
    )


def _original(tmp_path: Path) -> Path:
    """Un `original.docx` con cuerpo.

    Hace falta porque `generate_apa7_docx` busca el original en
    `salida.parent / "original.docx"` y, si no lo encuentra, FUERZA
    `cover_mode = "generate_apa7_template"`: sin archivo de entrada la portada
    UNI no llega a construirse y un test de logos mediria otra cosa.
    """
    tmp_path.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for i in range(8):
        d.add_paragraph(
            f"Parrafo {i} del cuerpo, con texto suficiente para que el gate de "
            "sanidad no se dispare por la portada que se reemplaza."
        )
    d.save(str(tmp_path / "original.docx"))
    return tmp_path


def generar(tmp_path: Path, portada: PortadaData, page_size: str = "carta") -> docx.Document:
    session = f"s-{abs(hash((portada.institution or '', page_size))) % 10**8}"
    carpeta = _original(tmp_path / session)
    salida = carpeta / f"{page_size}.docx"
    generate_apa7_docx(
        _modelo(session),
        salida,
        rules=APARuleSet(page_size=page_size),
        portada=portada,
    )
    return docx.Document(str(salida))


def con_institucion(nombre: str) -> PortadaData:
    return PortadaData(
        use_original_cover=False,
        cover_mode="generate_uni_cover",
        institution=nombre,
        title="Titulo",
        course="Asignatura",
    )


def con_logos(*logos: LogoPortada) -> PortadaData:
    return PortadaData(
        use_original_cover=False,
        cover_mode="generate_uni_cover",
        title="Titulo",
        course="Asignatura",
        logos=list(logos),
    )


def logo(asset: str, ancho_fraccion: float | None = None, institucion: str | None = None) -> LogoPortada:
    """Un logo pedido. Sin `ancho_fraccion` usa el DEFAULT del modelo, que es el
    numero que la portada UNI tenía antes de esta fase y el que hay que
    conservar."""
    if ancho_fraccion is None:
        return LogoPortada(asset=asset, institucion=institucion)
    return LogoPortada(asset=asset, ancho_fraccion=ancho_fraccion, institucion=institucion)


ASSETS = Path(__file__).parent.parent / "assets"


def _hash_de_un_asset(nombre: str) -> str:
    """La huella del archivo de `python/assets/`, o "" si no existe."""
    ruta_asset = ASSETS / nombre
    if not ruta_asset.exists():
        return ""
    return hashlib.sha1(ruta_asset.read_bytes()).hexdigest()


def _imagenes_del_docx(doc: docx.Document) -> list[tuple[str, str]]:
    """`(nombre_interno, huella)` de cada imagen del `.docx`.

    El nombre interno NO sirve para comparar: `python-docx` renombra la parte a
    `image1.png` al incorporarla, así que buscar "logo_unan.png" en el paquete no
    encuentra nunca nada. Lo que identifica a un logo es su CONTENIDO, y eso es
    lo que se compara: la pregunta real es si el documento lleva la imagen que
    pidió, no cómo se llama adentro.
    """
    salida = []
    vistas = set()
    for parte in doc.part.package.iter_parts():
        for rel in parte.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                destino = rel.target_part
            except Exception:
                continue
            clave = str(destino.partname)
            if clave in vistas:
                continue
            vistas.add(clave)
            salida.append((clave.rsplit("/", 1)[-1], hashlib.sha1(destino.blob).hexdigest()))
    return salida


def _usa_asset(doc: docx.Document, nombre: str) -> bool:
    """¿El `.docx` lleva ESTE logo? Por contenido, no por nombre de parte."""
    huella = _hash_de_un_asset(nombre)
    if not huella:
        return False
    return any(h == huella for _nombre, h in _imagenes_del_docx(doc))


def _assets_de_imagen(doc: docx.Document) -> list[str]:
    """Los NOMBRES de `python/assets/` que el documento contiene, para poder
    decir en el mensaje de fallo qué logo se coló."""
    hashes = {h for _n, h in _imagenes_del_docx(doc)}
    return sorted(
        p.name for p in ASSETS.glob("*.png") if hashlib.sha1(p.read_bytes()).hexdigest() in hashes
    )


def _ancho_de_imagen_cm(doc: docx.Document) -> float:
    """El ancho de la primera imagen, leído del XML del `.docx`."""
    ext = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    for parrafo in doc.paragraphs:
        for nodo in parrafo._element.iter(ext):
            return int(nodo.get("cx")) / EMU_POR_CM
    raise AssertionError("el documento no tiene ninguna imagen")


def _anchos_de_imagen_cm(doc: docx.Document) -> list[float]:
    ext = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    salida = []
    for parrafo in doc.paragraphs:
        for nodo in parrafo._element.iter(ext):
            salida.append(int(nodo.get("cx")) / EMU_POR_CM)
    return salida


# ── El defecto ───────────────────────────────────────────────────────────────

class TestElegirUnaInstitucionNoTraeElLogoDeOtra:
    def test_elegir_unan_no_pone_el_logo_de_uni(self, tmp_path):
        """Hoy `_resolve_logo_path()` no tiene argumento y siempre termina en
        `logo_uni.png`. Elegías UNAN y el `.docx` salía con el logo de la UNI, en
        silencio, porque `if LOGO_PATH.exists()` solo salta si el archivo no
        está."""
        doc = generar(tmp_path, con_institucion("UNAN-Managua"))
        assert _assets_de_imagen(doc), "no se coloco ninguna imagen en la portada"
        assert not _usa_asset(doc, "logo_uni.png"), (
            f"la portada de UNAN lleva el logo de la UNI: {_assets_de_imagen(doc)}"
        )

    def test_uni_sigue_llevando_su_propio_logo(self, tmp_path):
        """La otra mitad: la portada UNI no puede romperse en el arreglo."""
        doc = generar(tmp_path, con_institucion("UNI"))
        assert _usa_asset(doc, "logo_uni.png")

    def test_unan_si_lleva_el_suyo(self, tmp_path):
        doc = generar(tmp_path, con_institucion("UNAN-Managua"))
        assert _usa_asset(doc, "logo_unan.png")


class TestVariosLogos:
    def test_con_dos_logos_llegan_los_dos(self, tmp_path):
        doc = generar(
            tmp_path,
            con_logos(logo("logo_uni.png", 0.315), logo("logo_unan.png", 0.10)),
        )
        assert _usa_asset(doc, "logo_uni.png")
        assert _usa_asset(doc, "logo_unan.png")

    def test_cada_logo_mide_su_fraccion(self, tmp_path):
        doc = generar(
            tmp_path,
            con_logos(logo("logo_uni.png", 0.315), logo("logo_unan.png", 0.10)),
        )
        util_cm = portada_uni.ancho_util_mm("carta") / 10.0
        fracciones = sorted(round(w / util_cm, 2) for w in _anchos_de_imagen_cm(doc))
        assert fracciones == [0.10, 0.32]

    def test_un_asset_que_no_existe_no_inventa_una_imagen(self, tmp_path):
        """Un logo que se pidió y no llegó se avisa, no se sustituye en
        silencio por el de otro. Es la diferencia entre un documento incompleto y
        un documento equivocado."""
        doc = generar(tmp_path, con_logos(logo("logo_que_no_existe.png", 0.315)))
        assert not _usa_asset(doc, "logo_que_no_existe.png")
        assert not _usa_asset(doc, "logo_uni.png"), (
            "un asset inexistente no puede terminar poniendo el logo de la UNI"
        )


# ── La fracción del ancho útil ──────────────────────────────────────────────

class TestElLogoMideSuFraccionDelAnchoUtil:
    def test_el_logo_mide_su_fraccion_del_ancho_util_en_carta_y_en_a4(self):
        """El Review Focus #4. Con mm absolutos el logo se ve distinto en cada
        hoja, porque el ancho util de Carta y de A4 no es el mismo."""
        carta = generar(_tmp(), con_logos(logo("logo_uni.png", 0.315)), page_size="carta")
        a4 = generar(_tmp(), con_logos(logo("logo_uni.png", 0.315)), page_size="a4")
        util_carta = portada_uni.ancho_util_mm("carta") / 10.0
        util_a4 = portada_uni.ancho_util_mm("a4") / 10.0
        frac_carta = _ancho_de_imagen_cm(carta) / util_carta
        frac_a4 = _ancho_de_imagen_cm(a4) / util_a4
        assert frac_carta == pytest.approx(0.315, abs=0.01)
        assert frac_a4 == pytest.approx(0.315, abs=0.01)

    def test_el_ancho_util_de_verdad_es_16_51_cm_no_el_13_59_del_plan(self):
        """MEDIDO. El plan de la fase calcula sus numeros a partir de un ancho
        util de 13.59 cm, que sale de restar 40 mm por lado. El `.docx` real usa
        una pulgada: 165.1 mm de ancho util en carta y 159.2 mm en A4.

        De ese 13.59 cm salio tambien el default equivocado del logo, asi que
        esto esta aqui para que el error de las cuentas del plan no vuelva a
        entrar por la puerta del ancho util.
        """
        assert portada_uni.ancho_util_mm("carta") == pytest.approx(165.1, abs=0.05)
        assert portada_uni.ancho_util_mm("a4") == pytest.approx(159.2, abs=0.05)


class TestElLogoUniConservaSuTamano:
    """El logo UNI medía `Cm(5.2)` antes de esta fase y tiene que seguir
    midiendo eso. Un 16% del ancho util son 2.64 cm: la mitad, y el usuario lo
    reporta como "el logo que puso es super pequeño no se ve".

    La cuenta, que es la que fija el default:
        5.2 cm / 16.51 cm de ancho util = 0.315
    Los margenes son de una pulgada (`APARuleSet.margins_cm = 2.54`), no de
    40 mm, y por eso el ancho util es 16.51 cm y no 13.59 cm.
    """

    def test_el_logo_uni_conserva_su_tamano_anterior(self):
        """Un logo UNI y carta: 5.2 cm, que es lo que ponía el `Cm(5.2)`."""
        doc = generar(_tmp(), con_logos(logo("logo_uni.png")), page_size="carta")
        assert _ancho_de_imagen_cm(doc) == pytest.approx(5.2, abs=0.05)

    def test_el_logo_uni_es_la_misma_fraccion_en_a4(self):
        """En A4 el ancho ABSOLUTO cambia (el ancho util de A4 no es el de
        carta) y la FRACCIÓN no. Que se espere 0.31, y no 5.2 cm, es lo
        correcto: son hojas distintas."""
        doc = generar(_tmp(), con_logos(logo("logo_uni.png")), page_size="a4")
        util_a4_cm = portada_uni.ancho_util_mm("a4") / 10.0
        fraccion = _ancho_de_imagen_cm(doc) / util_a4_cm
        assert fraccion == pytest.approx(0.315, abs=0.01)
        # Y el absoluto es OTRO numero, no el mismo 5.2 cm repetido.
        assert _ancho_de_imagen_cm(doc) == pytest.approx(5.01, abs=0.05)


def _tmp():
    import tempfile

    return Path(tempfile.mkdtemp(prefix="logos-"))


# ── El modelo ────────────────────────────────────────────────────────────────

class TestElModeloDeLogos:
    def test_el_logo_default_es_una_fraccion_y_no_un_milimetro(self):
        # El plan declaraba `LogoPortada = {asset, ancho_mm, institucion}` y en
        # el mismo parrafo `ancho_fraccion: float = 0.16`. La fraccion es la que
        # funciona: un `ancho_mm` no puede ser el mismo en Carta y en A4. El
        # NUMERO del default si se corrigio: 0.16 salia de un ancho util de
        # 13.59 cm que el proyecto nunca tuvo.
        assert "ancho_fraccion" in LogoPortada.model_fields
        assert "ancho_mm" not in LogoPortada.model_fields
        # `asset` es obligatorio: un logo sin saber cual es no es un logo.
        # 0.315 es 5.2 cm sobre los 16.51 cm de ancho util de una carta: el
        # tamano que llevaba el `Cm(5.2)` de antes.
        assert LogoPortada(asset="logo_uni.png").ancho_fraccion == 0.315
        # Y el default del modelo y el fallback interno del generador son el
        # MISMO numero. Son dos copias de una sola cuenta, y si se separan el
        # logo cambia de tamano segun si la portada venga con lista de logos o
        # sin ella.
        assert LogoPortada(asset="logo_uni.png").ancho_fraccion == (
            portada_uni.FRACCION_DE_ANCHO_DEL_LOGO
        )
        with pytest.raises(Exception):
            LogoPortada()

    def test_la_fraccion_se_valida(self):
        # Una fraccion negativa o de mas de la hoja produce un `Cm()` que
        # python-docx acepta y Word no sabe dibujar. Que lo rechace el modelo.
        with pytest.raises(Exception):
            LogoPortada(asset="logo_uni.png", ancho_fraccion=0)
        with pytest.raises(Exception):
            LogoPortada(asset="logo_uni.png", ancho_fraccion=1.5)

    def test_portada_tiene_una_lista_de_logos(self):
        assert PortadaData().logos == []
        con_dos = PortadaData(logos=[logo("logo_uni.png"), logo("logo_unan.png", 0.1)])
        assert len(con_dos.logos) == 2


class TestSuprimirElLogo:
    """El hallazgo #6: la portada UNI ponía SIEMPRE un logo y el API no tenía cómo
    pedir una portada sin él. `mostrar_logo=False` lo omite por completo."""

    def test_mostrar_logo_false_no_inserta_imagen(self, tmp_path):
        portada = PortadaData(
            use_original_cover=False,
            cover_mode="generate_uni_cover",
            institution="UNI",
            title="Titulo",
            course="Asignatura",
            mostrar_logo=False,
        )
        doc = generar(tmp_path, portada)
        assert _imagenes_del_docx(doc) == [], (
            f"mostrar_logo=False no debe insertar el logo: {_assets_de_imagen(doc)}"
        )
