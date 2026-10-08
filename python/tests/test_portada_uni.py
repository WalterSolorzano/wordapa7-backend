"""Pruebas unitarias para el módulo de Portada UNI (portada_uni.py)."""

import sys
from pathlib import Path
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from modules.portada_uni import generate_uni_cover

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def test_generate_uni_cover_with_data():
    doc = Document()
    autores = [
        {"nombre": "Br. Walter Noel Solórzano Gaitán", "carnet": "2023-0432U"},
        {"nombre": "Br. Yireh Alejandro Beteta Tórrez", "carnet": "2023-0334U"},
    ]
    p_count = generate_uni_cover(
        doc,
        titulo="Caso Práctico: Herramienta Histograma",
        asignatura="Control Estadístico de la Calidad",
        autores=autores,
        tutor="Ing. María Auxiliadora",
        grupo="3T1 IND",
        departamento="Área de Conocimiento de Ingeniería y Afines",
        fecha="7 de septiembre del año 2026",
        lugar="Managua, Nicaragua",
    )
    assert p_count > 0

    # Verificar que todos los runs tengan <w:color w:val="000000"/>
    root = doc.element.body
    runs = root.xpath('.//w:r')
    assert len(runs) > 0

    for r in runs:
        colors = r.xpath('.//w:rPr/w:color/@w:val')
        if colors:
            assert colors[0] == "000000", f"Color no es negro puro: {colors[0]}"


def test_generate_uni_cover_editable_empty():
    doc = Document()
    # Generar portada vacía/editable
    p_count = generate_uni_cover(
        doc,
        titulo="",
        asignatura="",
        autores=[],
        tutor="",
        grupo="",
        fecha="",
    )
    assert p_count > 0

    # Verificar que contenga los marcadores por defecto
    full_text = "".join([p.text for p in doc.paragraphs])
    assert "Área de Conocimiento de Ingeniería y Afines" in full_text
    assert "Título del trabajo" in full_text
    assert "Elaborado por" in full_text
    assert "Managua, Nicaragua" in full_text


def test_tutor_sale_en_su_propia_columna_y_no_como_estudiante():
    """El docente debe salir SEPARADO de los estudiantes, en su propia columna
    (con su grupo), no mezclado en la lista de autores."""
    doc = Document()
    autores = [
        {"nombre": "Br. Estudiante Uno", "carnet": "2023-0001U"},
        {"nombre": "Br. Estudiante Dos", "carnet": "2023-0002U"},
    ]
    generate_uni_cover(
        doc,
        titulo="Caso Práctico",
        asignatura="Asignatura",
        autores=autores,
        tutor="Ing. Carlos Docente",
        grupo="3T1 IND",
        fecha="7 de septiembre del año 2026",
    )

    celdas = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    celda_tutor = [c for c in celdas if "Ing. Carlos Docente" in c]
    assert len(celda_tutor) == 1, f"el docente no tiene columna propia: {celdas}"
    # La columna del docente NO contiene estudiantes y sí trae el grupo.
    assert "Estudiante Uno" not in celda_tutor[0]
    assert "Estudiante Dos" not in celda_tutor[0]
    assert "3T1 IND" in celda_tutor[0]
