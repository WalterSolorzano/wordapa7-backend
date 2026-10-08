"""Tests de portadas maquetadas en tabla (propuesta 2).

Una portada puede vivir dentro de un ``w:tbl``. Hoy solo se leían textboxes y
párrafos, así que esas portadas se perdían. El extractor debe mirar SOLO la
zona de portada y exigir una señal de portada, para no confundir una tabla de
datos del cuerpo con una portada.
"""

import io

from docx import Document

from parsing.xml_deep_parser import (
    extract_cover_table_texts,
    extract_unique_textbox_pairs,
)


def _doc_with_table(rows, body_paragraphs=0):
    doc = Document()
    for _ in range(body_paragraphs):
        doc.add_paragraph("Contenido del cuerpo del trabajo con datos y análisis.")
    table = doc.add_table(rows=len(rows), cols=2)
    for r, (left, right) in enumerate(rows):
        table.cell(r, 0).text = left
        table.cell(r, 1).text = right
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Document(buf)


def test_tabla_de_portada_extrae_autores_y_carnets():
    doc = _doc_with_table(
        [
            ("Universidad Nacional de Ingeniería", "Facultad de Ingeniería"),
            ("Br. Ana López", "Carnet: 2022-0001"),
            ("Br. Luis Pérez", "Carnet: 2023-0002"),
        ]
    )
    lines = extract_cover_table_texts(doc)
    assert "Br. Ana López" in lines
    assert "Carnet: 2022-0001" in lines

    pairs = extract_unique_textbox_pairs(lines)
    ids = {p["id"] for p in pairs if p["role"] == "br."}
    assert {"2022-0001", "2023-0002"} <= ids


def test_tabla_del_cuerpo_no_se_confunde_con_portada():
    doc = _doc_with_table(
        [
            ("Variable", "Valor"),
            ("Costo", "100"),
            ("Tiempo", "5"),
        ],
        body_paragraphs=15,
    )
    assert extract_cover_table_texts(doc) == []


def test_tabla_sin_senal_de_portada_no_se_toma():
    doc = _doc_with_table(
        [
            ("Variable", "Valor"),
            ("Costo", "100"),
        ]
    )
    assert extract_cover_table_texts(doc) == []
