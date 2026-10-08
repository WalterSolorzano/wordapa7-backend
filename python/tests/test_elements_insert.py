"""FASE 3 — POST /api/elements/insert: dividir un párrafo en dos.

El párrafo físico se inserta en original.docx (misma convención de mapeo
por índice que apply_inplace) y el elemento en el modelo. Sin la inserción
física, apply_inplace desfasaría todos los elementos siguientes.
"""
import sys
import pathlib
import types

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient


def _client():
    from main import app
    return TestClient(app)


def _mk_doc(n=3):
    """DocumentModel real con n párrafos de cuerpo (ids e0..e{n-1})."""
    from models import DocumentModel, ElementModel, ElementType
    return DocumentModel(
        session_id="s1",
        elements=[
            ElementModel(
                id=f"e{i}", type=ElementType.PARAGRAPH,
                heading_level=1, text=f"texto {i}",
                is_cover_section=False,
            )
            for i in range(n)
        ],
    )


def _mk_session(tmp_path, monkeypatch, doc):
    """Sesión con original.docx real + estado mockeado en el router.

    sessions.py importa load/save por nombre: hay que parchear el namespace
    del router, no el módulo (misma técnica que test_update_element_table).
    """
    from routers import sessions as router
    from persistence import session_manager as sm
    sd = tmp_path / "sessions" / "s1"
    sd.mkdir(parents=True)
    (sd / "original.docx").write_bytes(b"PK\x03\x04fake")
    monkeypatch.setattr(router, "STORAGE_DIR", tmp_path)
    monkeypatch.setattr(router, "load_session_state", lambda sid, st: doc)
    monkeypatch.setattr(router, "save_session_state", lambda d, st: None)
    # save_session_snapshot se importa DENTRO de la función: parchear el módulo
    monkeypatch.setattr(sm, "save_session_snapshot", lambda d, st: None)
    return sd


def _write_real_docx(tmp_path, n_paras=3):
    """Crea un original.docx real con python-docx en la sesión."""
    from docx import Document
    d = Document()
    for i in range(n_paras):
        d.add_paragraph(f"texto {i}")
    d.save(str(tmp_path / "sessions" / "s1" / "original.docx"))


def test_insert_404_sesion_desconocida():
    r = _client().post("/api/elements/insert", json={
        "session_id": "nope", "after_element_id": "e0",
        "new_element_id": "n1", "text": "hola",
    })
    assert r.status_code == 404


def test_insert_404_elemento_desconocido(tmp_path, monkeypatch):
    _mk_session(tmp_path, monkeypatch, _mk_doc())
    r = _client().post("/api/elements/insert", json={
        "session_id": "s1", "after_element_id": "zzz",
        "new_element_id": "n1", "text": "hola",
    })
    assert r.status_code == 404


def test_insert_divide_parrafo_fisico_y_modelo(tmp_path, monkeypatch):
    doc = _mk_doc(3)
    _mk_session(tmp_path, monkeypatch, doc)
    _write_real_docx(tmp_path, n_paras=3)
    r = _client().post("/api/elements/insert", json={
        "session_id": "s1", "after_element_id": "e1",
        "new_element_id": "n1", "text": "mitad nueva",
    })
    assert r.status_code == 200
    data = r.json()
    ids = [e["id"] for e in data["elements"]]
    assert ids == ["e0", "e1", "n1", "e2"]
    texts = [e["text"] for e in data["elements"]]
    assert texts[2] == "mitad nueva"
    # El párrafo físico existe en el docx (4 párrafos)
    from docx import Document
    d = Document(str(tmp_path / "sessions" / "s1" / "original.docx"))
    assert len(d.paragraphs) == 4
    assert d.paragraphs[2].text == "mitad nueva"


def test_insert_no_toca_portada(tmp_path, monkeypatch):
    """Un párrafo de portada (is_cover_section) no puede ser dividido."""
    doc = _mk_doc(2)
    doc.elements[0].is_cover_section = True
    _mk_session(tmp_path, monkeypatch, doc)
    _write_real_docx(tmp_path, n_paras=2)
    r = _client().post("/api/elements/insert", json={
        "session_id": "s1", "after_element_id": "e0",
        "new_element_id": "n1", "text": "hola",
    })
    assert r.status_code == 400
