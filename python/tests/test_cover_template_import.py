from pathlib import Path

from docx import Document

from modules.cover_designer import detect_cover_fields_from_docx


def test_detect_cover_fields_from_docx_extracts_editable_metadata(tmp_path: Path):
    doc = Document()
    doc.add_paragraph('UNIVERSIDAD NACIONAL DE INGENIERÍA')
    doc.add_paragraph('Tema: Impacto de la automatización en la educación')
    doc.add_paragraph('Br. Ana Pérez | Carnet: 2023-1234')
    doc.add_paragraph('Docente: Ing. Carlos Rojas')
    doc.add_paragraph('Fecha: 12 de marzo de 2026')

    source = tmp_path / 'cover_example.docx'
    doc.save(source)

    result = detect_cover_fields_from_docx(source)

    assert result['detected'] is True
    assert result['fields']['title'] == 'Impacto de la automatización en la educación'
    assert 'Ana Pérez' in (result['fields'].get('author') or '')
    assert 'Carlos Rojas' in (result['fields'].get('instructor') or '')
    assert '2026' in (result['fields'].get('date') or '')
