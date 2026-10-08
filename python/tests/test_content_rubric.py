"""Tests del Traductor de Rúbrica (motor oculto, 0 tokens)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from docx import Document as DocxDocument  # python-docx

from content.rubric import build_rubric_report, evaluate_rubric, parse_rubric
from models import DocumentModel, ElementModel, ElementType, ReferenciaModel


def _make_xlsx(path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_parse_rubric_xlsx(tmp_path):
    path = _make_xlsx(
        tmp_path / "rubrica.xlsx",
        [
            ["Criterio", "Peso", "Descripción"],
            ["Portada", "10", "Portada institucional"],
            ["Introducción", "20", "Presenta el tema"],
            ["Referencias", "30", "APA 7"],
            ["Ortografía", "40", "Sin errores"],
        ],
    )
    criteria = parse_rubric(path)
    assert [c["name"] for c in criteria] == ["Portada", "Introducción", "Referencias", "Ortografía"]
    assert abs(sum(c["weight"] for c in criteria) - 100) < 0.01
    assert criteria[0]["description"] == "Portada institucional"


def test_parse_rubric_normalizes_weights(tmp_path):
    path = _make_xlsx(
        tmp_path / "r.xlsx",
        [
            ["Criterio", "Peso"],
            ["A", "1"],
            ["B", "2"],
            ["C", "3"],
        ],
    )
    criteria = parse_rubric(path)
    weights = [c["weight"] for c in criteria]
    assert abs(sum(weights) - 100) < 0.5
    assert weights == [pytest.approx(16.67), pytest.approx(33.33), pytest.approx(50.0)]


def test_parse_rubric_docx_table(tmp_path):
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Criterio"
    table.rows[0].cells[1].text = "Peso"
    for name, weight in [("Metodología", "50"), ("Resultados", "50")]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = weight
    path = tmp_path / "rubrica.docx"
    doc.save(path)

    criteria = parse_rubric(path)
    assert [c["name"] for c in criteria] == ["Metodología", "Resultados"]
    assert abs(sum(c["weight"] for c in criteria) - 100) < 0.01


def _doc_with(headings=(), paragraphs=(), references=0, tables=0):
    elements = []
    for h in headings:
        elements.append(
            ElementModel(id=f"h-{h}", type=ElementType.HEADING, text=h, heading_level=1)
        )
    for p in paragraphs:
        elements.append(ElementModel(id=f"p-{len(elements)}", type=ElementType.PARAGRAPH, text=p))
    for i in range(tables):
        elements.append(ElementModel(id=f"t-{i}", type=ElementType.TABLE))
    doc = DocumentModel(session_id="s1", elements=elements)
    doc.referencias = [
        ReferenciaModel(id=f"ref-{i}", raw_text=f"Obra {i}.", formatted_apa=f"Obra {i}.")
        for i in range(references)
    ]
    return doc


def test_evaluate_deterministic():
    doc = _doc_with(headings=["Introducción"], paragraphs=["El objetivo general es X."])
    criteria = [
        {"name": "Introducción", "weight": 20, "description": ""},
        {"name": "Referencias", "weight": 30, "description": ""},
        {"name": "Ortografía", "weight": 50, "description": ""},
    ]
    results, summary = evaluate_rubric(criteria, doc)
    by_name = {r["name"]: r for r in results}

    assert by_name["Introducción"]["score"] == 100
    assert by_name["Introducción"]["status"] == "cumple"
    assert by_name["Referencias"]["score"] == 0
    assert by_name["Referencias"]["status"] == "no_cumple"
    assert by_name["Ortografía"]["score"] is None
    assert by_name["Ortografía"]["status"] == "requiere_ia"

    # Solo se promedian los criterios medibles: (100*20 + 0*30) / 50 = 40
    assert summary["score"] == 40.0
    assert summary["coverage_pct"] == 50.0
    assert summary["scored_criteria"] == 2


def test_build_rubric_report_end_to_end(tmp_path):
    rubric = _make_xlsx(
        tmp_path / "rubrica.xlsx",
        [
            ["Criterio", "Peso"],
            ["Portada", "20"],
            ["Introducción", "40"],
            ["Ortografía", "40"],
        ],
    )
    src = DocxDocument()
    src.add_heading("Introducción", level=1)
    src.add_paragraph("Texto de prueba.")
    docx_path = tmp_path / "tesis.docx"
    src.save(docx_path)

    report = build_rubric_report(rubric, docx_path, storage_dir=tmp_path)
    assert report["rubric"]["criteria_count"] == 3
    assert report["summary"]["total_criteria"] == 3
    assert [c["name"] for c in report["criteria"]] == ["Portada", "Introducción", "Ortografía"]
    json.dumps(report)  # serializable
