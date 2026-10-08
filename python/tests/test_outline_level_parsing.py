"""OutLvl real de Word: el parser debe leer `w:outlineLvl` del DOCX.

Hoy `ElementModel.heading_level` lo decide el pre-clasificador por formato, y
`StyleFingerprint.outline_level` estaba hardcodeado a 9.0 en
`clustering_classifier.py`. Eso ignora un documento que marca sus títulos solo
con el outline de Word (sin estilo Heading ni negrita). Estos tests fijan que
el nivel crudo del outline llega al modelo.
"""

from docx import Document

from parsing.docx_parser import parse_docx_bytes

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _guardar(doc, tmp_path, nombre):
    path = tmp_path / nombre
    doc.save(path)
    return path.read_bytes()


def _parrafo_con_outline(doc, texto, nivel):
    p = doc.add_paragraph(texto)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(f"{W}outlineLvl", {f"{W}val": str(nivel)}))
    return p


def test_lee_outline_lvl_cero_como_valor_crudo(tmp_path):
    doc = Document()
    _parrafo_con_outline(doc, "Resultados", 0)
    data = _guardar(doc, tmp_path, "outline.docx")

    modelo = parse_docx_bytes(data, "outline.docx", "s1", tmp_path)
    p = next(e for e in modelo.elements if (e.text or "").strip() == "Resultados")
    assert p.outline_level == 0


def test_parrafo_sin_outline_tiene_none(tmp_path):
    doc = Document()
    doc.add_paragraph("Texto normal")
    data = _guardar(doc, tmp_path, "sin.docx")

    modelo = parse_docx_bytes(data, "sin.docx", "s1", tmp_path)
    p = next(e for e in modelo.elements if (e.text or "").strip() == "Texto normal")
    assert p.outline_level is None


def test_outline_cero_clasifica_como_h1_sin_formato(tmp_path):
    """Un titulo marcado SOLO por outline (sin negrita ni estilo) es H1.

    Este es el caso que hoy cae a parrafo: el autor uso la vista de esquema de
    Word y no el estilo Heading, asi que no hay formato que mirar. El outline
    real (w:outlineLvl=0) es la senal que lo delata.
    """
    doc = Document()
    _parrafo_con_outline(doc, "Resultados de la investigacion", 0)
    data = _guardar(doc, tmp_path, "o0.docx")

    modelo = parse_docx_bytes(data, "o0.docx", "s1", tmp_path)
    p = next(e for e in modelo.elements if (e.text or "").strip() == "Resultados de la investigacion")
    assert p.type == "heading"
    assert p.heading_level == 1


def test_outline_uno_clasifica_como_h2_sin_formato(tmp_path):
    doc = Document()
    _parrafo_con_outline(doc, "Antecedentes teoricos", 1)
    data = _guardar(doc, tmp_path, "o1.docx")

    modelo = parse_docx_bytes(data, "o1.docx", "s1", tmp_path)
    p = next(e for e in modelo.elements if (e.text or "").strip() == "Antecedentes teoricos")
    assert p.type == "heading"
    assert p.heading_level == 2
