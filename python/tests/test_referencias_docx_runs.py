"""La sección de Referencias del DOCX escribe un run por segmento (cursiva)."""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import docx

from models import APARuleSet, ReferenciaModel
from modules.referencias_module import format_apa_referencias_section

RULES = APARuleSet(font_family="Times New Roman", font_size_pt=12)


def _runs_of_reference(doc):
    for p in doc.paragraphs:
        if p.text.startswith("Hirano"):
            return p.runs
    return []


def test_docx_escribe_run_cursivo():
    ref = ReferenciaModel(
        id="r1", authors=["Hirano, H."], year="1995",
        title="5 Pillars of the Visual Workplace", source="Productivity Press",
        tipo="libro",
    )
    doc = docx.Document()
    format_apa_referencias_section(doc, [ref], RULES)
    runs = _runs_of_reference(doc)
    assert any(r.italic and "5 Pillars" in r.text for r in runs)
    assert any((not r.italic) and "Productivity Press" in r.text for r in runs)
    assert "[OIT]" not in "".join(r.text for r in runs)


def test_docx_corporativo_sin_siglas():
    ref = ReferenciaModel(
        id="r2", authors=["Instituto Nicaragüense de Energía [INE]"], year="2026",
        title="Informe", source="INE", tipo="informe",
    )
    doc = docx.Document()
    format_apa_referencias_section(doc, [ref], RULES)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Instituto Nicaragüense de Energía" in texto
    assert "[INE]" not in texto


def test_docx_usa_formatted_apa_sin_campos():
    # Regresión: una referencia que solo trae formatted_apa no debe desaparecer.
    ref = ReferenciaModel(id="r3", formatted_apa="Autor, A. (2019). Texto que debe salir.")
    doc = docx.Document()
    format_apa_referencias_section(doc, [ref], RULES)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Texto que debe salir" in texto
