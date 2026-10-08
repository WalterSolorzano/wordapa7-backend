"""Tests de integracion de POST /api/spec (generacion mockeada salvo E2E)."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
from fastapi.testclient import TestClient


@pytest.fixture(scope="session", autouse=True)
def _ensure_template():
    """Crea la plantilla APA base una vez (en tests no corre el lifespan)."""
    from config import get_apa7_template_path
    p = get_apa7_template_path()
    if not p.exists():
        from create_template import create_apa7_template
        create_apa7_template(p)


@pytest.fixture
def client():
    """TestClient sin context manager (evita lifespan/COM)."""
    from main import app
    return TestClient(app)


def _spec(**over):
    base = {
        "spec_version": "1",
        "elements": [{"type": "heading", "level": 1, "text": "1. Marco"},
                     {"type": "paragraph", "text": "Texto de prueba."}],
    }
    base.update(over)
    return base


def test_unknown_version_422(client):
    r = client.post("/api/spec", json=_spec(spec_version="9"))
    assert r.status_code == 422


def test_unknown_preset_404_with_available(client):
    r = client.post("/api/spec", json=_spec(presets={"table": "no_existe"}))
    assert r.status_code == 404
    body = r.json()["detail"]
    assert "tabla_apa_generica" in body["available"]


def test_unknown_cover_404_with_available(client):
    r = client.post("/api/spec", json=_spec(cover={"template": "portada_fantasma"}))
    assert r.status_code == 404
    assert "available" in r.json()["detail"]


def test_data_uri_422(client):
    r = client.post("/api/spec", json=_spec(elements=[
        {"type": "figure", "image": "data:image/png;base64,AA",
         "caption": "Figura 1", "title": "T"}]))
    assert r.status_code == 422


def test_happy_path_200(client, monkeypatch):
    """Generacion mockeada: valida orquestacion, sesion y respuesta."""
    import docx
    import main

    async def fake_generate(req):
        out = Path(main.STORAGE_DIR) / "sessions" / req.session_id
        out.mkdir(parents=True, exist_ok=True)
        d = docx.Document()
        d.add_paragraph("1. Marco")
        d.add_paragraph("Texto de prueba.")
        d.save(str(out / "APA7_test.docx"))
        return {"download_url": f"/api/download/{req.session_id}",
                "filename": "APA7_test.docx"}

    monkeypatch.setattr(main, "generate_docx", fake_generate)
    r = client.post("/api/spec", json=_spec(presets={"layout": "layout_uni"}))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["session_id"] and body["download_url"].endswith(body["session_id"])
    assert body["filename"] == "APA7_test.docx"
    assert isinstance(body["warnings"], list)


def test_missing_image_warning_200(client, monkeypatch):
    import docx
    import main

    async def fake_generate(req):
        out = Path(main.STORAGE_DIR) / "sessions" / req.session_id
        out.mkdir(parents=True, exist_ok=True)
        d = docx.Document()
        d.add_paragraph("x")   # texto del spec: satisface el sanity gate
        d.save(str(out / "APA7_w.docx"))
        return {"download_url": f"/api/download/{req.session_id}",
                "filename": "APA7_w.docx"}

    monkeypatch.setattr(main, "generate_docx", fake_generate)
    r = client.post("/api/spec", json=_spec(elements=[
        {"type": "paragraph", "text": "x"},
        {"type": "figure", "image": "C:/no_existe/a.png",
         "caption": "Figura 1", "title": "T"}]))
    assert r.status_code == 200
    assert any("no_existe" in w for w in r.json()["warnings"])


def test_wrong_type_preset_422(client):
    """Preset de tipo distinto al pedido -> 422 autocorregible."""
    r = client.post("/api/spec", json=_spec(presets={"table": "layout_uni"}))
    assert r.status_code == 422


def test_table_element_unknown_preset_404(client):
    """Preset inexistente en TableElement.preset -> 404 con available."""
    r = client.post("/api/spec", json=_spec(elements=[
        {"type": "table", "caption": "Tabla 1", "title": "T",
         "columns": ["a"], "rows": [["1"]], "preset": "no_tal"}]))
    assert r.status_code == 404
    assert "tabla_apa_generica" in r.json()["detail"]["available"]


def test_table_element_wrong_type_preset_422(client):
    """Preset de tipo != table en TableElement.preset -> 422, no silencio."""
    r = client.post("/api/spec", json=_spec(elements=[
        {"type": "table", "caption": "Tabla 1", "title": "T",
         "columns": ["a"], "rows": [["1"]], "preset": "layout_uni"}]))
    assert r.status_code == 422


def test_equipment_card_data_uri_422(client):
    r = client.post("/api/spec", json=_spec(elements=[
        {"type": "equipment_card", "number": "A1", "title": "Eq",
         "image": "data:image/png;base64,AA", "specs": {}}]))
    assert r.status_code == 422


def test_cover_scratch_pasa_portada_a_generate(client, monkeypatch):
    """cover SIN template -> PortadaData scratch (use_original_cover=False)."""
    import docx
    import main

    captured = {}

    async def fake_generate(req):
        captured["portada"] = req.portada
        captured["meta"] = req.meta
        out = Path(main.STORAGE_DIR) / "sessions" / req.session_id
        out.mkdir(parents=True, exist_ok=True)
        d = docx.Document()
        d.add_paragraph("1. Marco")
        d.add_paragraph("Texto de prueba.")
        d.save(str(out / "APA7_test.docx"))
        return {"download_url": f"/api/download/{req.session_id}",
                "filename": "APA7_test.docx"}

    monkeypatch.setattr(main, "generate_docx", fake_generate)
    r = client.post("/api/spec", json=_spec(cover={
        "title": "Balance energetico", "author": "Walter Solorzano",
        "institution": "UNI", "course": "Tecnologia y Medio Ambiente",
        "instructor": "Ing. Eva Mairena", "date": "octubre 2026"}))
    assert r.status_code == 200, r.text
    p = captured.get("portada")
    assert p is not None, "cover scratch no llego a generate_docx"
    assert p.use_original_cover is False, "scratch debe pedir portada sintetica"
    assert p.title == "Balance energetico"
    assert p.institution == "UNI"
    assert p.course == "Tecnologia y Medio Ambiente"
    assert p.date == "octubre 2026"

    # El autor y el profesor NO viajan en la portada: viajan en los metadatos
    # del documento. Con una plantilla cargada la portada original se conserva
    # intacta, asi que un autor guardado en `portada` no tendria de donde salir
    # y el `.docx` saldria sin el. Ver el motivo en `models.py`.
    m = captured.get("meta")
    assert m is not None, "el acta no llego a generate_docx"
    assert m.autor == "Walter Solorzano"
    assert m.profesor_asesor == ["Ing. Eva Mairena"]


def test_cover_vacio_sin_template_422(client):
    """cover sin template y sin un solo dato -> 422 (nada que renderizar)."""
    r = client.post("/api/spec", json=_spec(cover={}))
    assert r.status_code == 422, (
        "cover vacio debe rechazarse, no generarse sin portada")


# ==-==- Ejemplos OpenAPI copiables por agentes (mejora #5, pieza 4) ==-==


def _spec_examples(openapi: dict) -> list[dict]:
    """Ejemplos del body de POST /api/spec, en cualquier colocacion
    (media-type `examples` o `schema.examples`; cambia entre versiones
    de FastAPI)."""
    media = (openapi["paths"]["/api/spec"]["post"]
             ["requestBody"]["content"]["application/json"])
    found: list[dict] = []

    def _walk(node):
        if isinstance(node, dict):
            if "spec_version" in node:
                found.append(node)
            for v in node.values():
                _walk(v)
        elif isinstance(node, list):
            for v in node:
                _walk(v)

    _walk(media)
    return found


def test_openapi_ejemplos_spec_validos_y_usables(client):
    """POST /api/spec publica >=2 ejemplos que validan contra el DSL
    y cuyos presets/portada referenciados existen de verdad."""
    import json as _json

    from config import STORAGE_DIR
    from modules.cover_designer import list_cover_templates
    from preset_store import PresetNotFound, get_preset
    from spec_dsl import SpecDocument

    r = client.get("/openapi.json")
    assert r.status_code == 200
    openapi = r.json()

    blob = _json.dumps(openapi, ensure_ascii=False).lower()
    assert "completo" in blob, "resumen del ejemplo completo ausente"
    assert "minimo" in blob, "resumen del ejemplo minimo ausente"

    examples = _spec_examples(openapi)
    assert len(examples) >= 2, (
        f"se esperaban >=2 ejemplos de spec, hay {len(examples)}")

    for ex in examples:
        SpecDocument.model_validate(ex)   # lanza ValidationError si falla

    completo = next((e for e in examples if e.get("cover")), None)
    assert completo, "el ejemplo completo debe incluir bloque cover"
    for pname in completo.get("presets", {}).values():
        try:
            get_preset(pname, STORAGE_DIR)
        except PresetNotFound:
            pytest.fail(f"el ejemplo completo apunta al preset inexistente "
                        f"'{pname}'")
    cover_names = [t.name for t in list_cover_templates(STORAGE_DIR)]
    assert completo["cover"]["template"] in cover_names, (
        f"template '{completo['cover']['template']}' inexistente")


def test_openapi_ejemplo_type_cover(client):
    """El query param type de GET /api/presets publica el ejemplo 'cover'."""
    import json as _json

    r = client.get("/openapi.json")
    params = r.json()["paths"]["/api/presets"]["get"]["parameters"]
    t = next((p for p in params if p.get("name") == "type"), None)
    assert t is not None, "parametro type ausente del openapi"
    dumped = _json.dumps(t)
    # El VALOR cover como ejemplo, no la palabra en la descripcion
    assert ('"value": "cover"' in dumped
            or '"example": "cover"' in dumped
            or t.get("schema", {}).get("example") == "cover"), (
        f"el parametro type no publica ejemplo con valor cover: {dumped}")
