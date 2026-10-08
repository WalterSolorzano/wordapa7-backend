"""
WordAPA7 — Motor de Estilos APA 7

Aplica reglas estrictas APA 7 a nivel de tipografia, parrafos y titulos (Nivel 1 a 5).
Normaliza la fuente en todo el documento para evitar mezclas indeseadas.
"""

from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from models import APARuleSet


def set_run_font(run, font_family: str, font_size_pt: float) -> None:
    """
    Aplica fuente y tamaño de forma robusta a nivel de run, incluyendo
    atributos rFonts (ascii, hAnsi, eastAsia, cs) que python-docx no setea
    por defecto. Elimina explícitamente atributos de tema (w:asciiTheme, w:hAnsiTheme,
    etc.) para evitar que Microsoft Word con temas activos (Calibri/Aptos)
    sobreescriba la tipografía APA 7 (F-10).
    """
    run.font.name = font_family
    run.font.size = Pt(font_size_pt)
    rPr = run._element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)

    # F-10: Eliminar cualquier tema que invalide la fuente explícita
    for theme_attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        theme_qn = qn(f'w:{theme_attr}')
        if theme_qn in rFonts.attrib:
            del rFonts.attrib[theme_qn]

    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), font_family)
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz')
        rPr.append(sz)
    sz.set(qn('w:val'), str(int(font_size_pt * 2)))
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), str(int(font_size_pt * 2)))


# ── TAMAÑO DE HOJA ────────────────────────────────────────────────────────────
# En milímetros, que es como las dos medidas se escriben sin ambigüedad: 8.5" son
# 215.9 mm y 11" son 279.4 mm. Pulgadas y milímetros no dividen igual, y un
# redondeo a 8.49" es una hoja que Word no reconoce como Carta.
#
# `python-docx` no tiene un nombre de hoja: el tamaño ES `section.page_width` y
# `section.page_height`, en EMU. Por eso esta tabla no puede vivir en el modelo
# y por eso el generador tiene que llamar a esta función.
TAMANOS_DE_PAGINA_MM: dict[str, tuple[float, float]] = {
    "carta": (215.9, 279.4),
    "a4": (210.0, 297.0),
}

EMU_POR_MM = 36000


def aplicar_tamano_pagina(section, page_size: str, landscape: bool = False) -> None:
    """Fija `page_width`/`page_height` de UNA sección al tamaño pedido.

    En landscape se escriben el alto y el ancho cruzados, y la orientación se
    declara en el `w:pgSz`: si se escribieran cruzados sin cambiarla, Word muestra
    una hoja apaisada con el flag de portrait, que es un caso distinto.
    """
    ancho_mm, alto_mm = TAMANOS_DE_PAGINA_MM.get(str(page_size or "").lower(), TAMANOS_DE_PAGINA_MM["carta"])
    if landscape:
        ancho_mm, alto_mm = alto_mm, ancho_mm
    section.page_width = int(round(ancho_mm * EMU_POR_MM))
    section.page_height = int(round(alto_mm * EMU_POR_MM))
    try:
        from docx.enum.section import WD_ORIENT
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    except Exception:
        pass


# ── IDIOMA DEL DOCUMENTO ──────────────────────────────────────────────────────
# `w:lang` es lo que Word usa para decidir qué corrector de ortografía aplicar.
# Sin él —o con el `en-US` que trae la plantilla de Word— un texto en español
# sale con cada palabra subrayada, y quien lo ve concludes que el documento está
# mal escrito en vez de que el archivo no lo declara.
#
# Se escribe en TRES lugares a propósito, y los tres hacen falta:
#   1. `docDefaults/rPr` — el valor por omisión de todo el documento.
#   2. el estilo `Normal` — que es donde Word mira primero.
#   3. cada run — porque un documento que viene de Word trae `w:lang` puesto en
#      el run, y ese gana sobre los defaults. Sin el punto 3, elegir el idioma
#      no cambiaría nada en la mitad de los documentos reales.
#
# `w:lang` es un atributo de corrección, no de apariencia: reescribirlo no
# cambia ni una coma de lo que se ve.


def _escribir_lang(rPr, language: str) -> None:
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), language)


def aplicar_idioma_documento(
    doc: docx.Document,
    language: str,
    skip_body_paragraphs: int = 0,
) -> None:
    """Declara en el `.docx` en qué idioma está escrito. Ver la nota de arriba.

    `skip_body_paragraphs` es el piso del cuerpo: los párrafos anteriores son la
    portada, y la portada es zona protegida (`AGENTS.md`), así que ahí no se
    escribe ni un byte. El nombre es el de `normalize_all_fonts` a propósito: las
    dos funciones caminan el documento por la misma frontera.
    """
    if not language:
        return

    # 1 y 2: defaults y estilo Normal. Son parte de la hoja, no de un párrafo, así
    # que no tocan la portada.
    try:
        styles_element = doc.styles._element
        doc_defaults = styles_element.find(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docDefaults"
        )
        if doc_defaults is not None:
            rpr = doc_defaults.find(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr"
            )
            if rpr is not None:
                _escribir_lang(rpr, language)
        for style in styles_element.iter(qn("w:style")):
            name_el = style.find(qn("w:name"))
            name_val = name_el.attrib.get(qn("w:val")) if name_el is not None else ""
            if name_val != "Normal":
                continue
            # Solo si el estilo YA tiene `rPr`. Crear uno donde no lo hay deja un
            # `<w:rPr/>` vacío en la hoja, que es un cambio que no es el idioma y
            # que la prueba de contrato no debería tener que aprender a ignorar.
            rpr_estilo = style.find(qn("w:rPr"))
            if rpr_estilo is not None:
                _escribir_lang(rpr_estilo, language)
    except Exception:
        pass

    # 3: los runs, que es donde gana el `w:lang` heredado del original.
    for para in doc.paragraphs[max(skip_body_paragraphs, 0):]:
        _escribir_lang_en_runs(para._element, language)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _escribir_lang_en_runs(para._element, language)


def _escribir_lang_en_runs(elemento, language: str) -> None:
    for run in elemento.iter(qn("w:r")):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            run.insert(0, rpr)
        _escribir_lang(rpr, language)


def apply_page_setup(doc: docx.Document, rules: APARuleSet, preserve_landscape: bool = True) -> None:
    """
    Aplica el TAMAÑO DE HOJA y los márgenes APA 7 (2.54 cm / 1 in).

    EL TAMAÑO DE HOJA ES DE `rules.page_size`, y antes no lo era: cada ruta
    hacía lo que podía (una copiaba el original, otra escribía Letter a mano) y el
    lienzo usaba un token fijo de A4. El resultado era que la pantalla y el
    archivo no tenían el mismo papel, y nadie lo notaba porque los dos tamaños
    parecían razonables. Ahora hay un solo lugar que decide, y es este.

    Las secciones en landscape conservan su orientación: se les intercambian
    ancho y alto para que el mismo contenido quepa girado, y sus márgenes
    originales quedan como estaban.
    """
    margin_inches = rules.margins_cm / 2.54
    page_size = getattr(rules, "page_size", "carta") or "carta"

    for section in doc.sections:
        is_landscape: bool = False
        try:
            pg_sz = section._sectPr.find(qn("w:pgSz"))
            if pg_sz is not None:
                orient = pg_sz.attrib.get(qn("w:orient"))
                if orient == "landscape":
                    is_landscape = True
        except Exception:
            pass

        aplicar_tamano_pagina(section, page_size, landscape=is_landscape)

        # En portrait: margenes APA 7. En landscape: conservar todo (orientacion y margenes).
        if not (preserve_landscape and is_landscape):
            section.top_margin = Inches(margin_inches)
            section.bottom_margin = Inches(margin_inches)
            section.left_margin = Inches(margin_inches)
            section.right_margin = Inches(margin_inches)


def normalize_all_fonts(
    doc: docx.Document,
    rules: Optional[APARuleSet] = None,
    target_font: Optional[str] = None,
    target_size_pt: Optional[int] = None,
    skip_body_paragraphs: int = 0,
) -> None:
    """
    Normaliza TODAS las fuentes en el documento a la fuente APA.

    REGLA DURA — se aplica a:
    - Parrafos normales (salvo los primeros `skip_body_paragraphs`, que son la portada)
    - Runs dentro de parrafos (fuentes mixtas pegadas de internet)
    - Texto dentro de celdas de tabla (ALL cells)
    - Encabezados y pies de pagina

    EXCEPCION: Preserva negrita/italica intencional (solo limpia font_name y font_size)
    NO AFECTA: Ecuaciones OMML, objetos OLE, texto en imagenes
    """
    target_font_pt = Pt(target_size_pt)

    # Normalizar parrafos del cuerpo (saltando la portada si se indica)
    for para in doc.paragraphs[skip_body_paragraphs:]:
        _normalize_paragraph_runs(para, target_font, target_font_pt)

    # Normalizar celdas de tabla
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _normalize_paragraph_runs(para, target_font, target_font_pt)

    # Normalizar encabezados y pies de pagina
    for section in doc.sections:
        # Header
        for para in section.header.paragraphs:
            _normalize_paragraph_runs(para, target_font, target_font_pt)
        # Footer
        for para in section.footer.paragraphs:
            _normalize_paragraph_runs(para, target_font, target_font_pt)

    # Normalizar estilos de la plantilla directamente en el XML de styles.xml
    if rules is None:
        rules = APARuleSet()
    update_docx_styles_xml(doc, rules)


def normalize_cover_font_name(
    doc: docx.Document,
    font_family: str,
    paragraph_count: int,
) -> None:
    """
    Cambia SOLO el nombre de la fuente en los párrafos de portada a la fuente
    APA (Times New Roman) pero RESPETA el tamaño original de cada run.
    Así la portada conserva su diseño y jerarquía visual, pero con la
    tipografía correcta de la norma.
    """
    if paragraph_count <= 0:
        return
    for para in doc.paragraphs[:paragraph_count]:
        for run in para.runs:
            run.font.name = font_family
            # rFonts OOXML para asegurar que Word no herede la fuente del estilo
            rPr = run._element.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                run._element.insert(0, rPr)
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.insert(0, rFonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(attr), font_family)


def update_docx_styles_xml(doc: docx.Document, rules: APARuleSet) -> None:
    """
    Modifica directamente las definiciones XML en styles.xml de la plantilla/documento:
    - Forzar fuente Times New Roman (12pt / sz=24)
    - Color negro (000000)
    - Interlineado doble (w:line="480", lineRule="auto")
    - Estilos Heading 1 a 5 y Normal / List Bullet
    """
    font_name = rules.font_family
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    try:
        styles_element = doc.styles._element

        # 1. Modificar docDefaults
        docDefaults = styles_element.find(f".//{{{w_ns}}}docDefaults")
        if docDefaults is not None:
            rPr = docDefaults.find(f".//{{{w_ns}}}rPr")
            if rPr is not None:
                rFonts = rPr.find(f"{{{w_ns}}}rFonts")
                if rFonts is None:
                    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
                    rPr.append(rFonts)
                rFonts.set(qn("w:ascii"), font_name)
                rFonts.set(qn("w:hAnsi"), font_name)

                sz = rPr.find(f"{{{w_ns}}}sz")
                if sz is None:
                    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="24"/>')
                    rPr.append(sz)
                else:
                    sz.set(qn("w:val"), "24")

            pPr = docDefaults.find(f".//{{{w_ns}}}pPr")
            if pPr is not None:
                spacing = pPr.find(f"{{{w_ns}}}spacing")
                if spacing is None:
                    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:line="480" w:lineRule="auto" w:before="0" w:after="0"/>')
                    pPr.append(spacing)
                else:
                    spacing.set(qn("w:line"), "480")
                    spacing.set(qn("w:lineRule"), "auto")

        # 2. Actualizar cada estilo individual
        heading_configs = {
            "Normal": {"color": "000000", "sz": "24", "line": "480", "bold": False, "italic": False},
            "Heading 1": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": False, "jc": "center"},
            "Heading 2": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": False, "jc": "left"},
            "Heading 3": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": True, "jc": "left"},
            "Heading 4": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": False, "jc": "left"},
            "Heading 5": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": True, "jc": "left"},
            "Título 1": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": False, "jc": "center"},
            "Título 2": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": False, "jc": "left"},
            "Título 3": {"color": "000000", "sz": "24", "line": "480", "bold": True, "italic": True, "jc": "left"},
            "List Bullet": {"color": "000000", "sz": "24", "line": "480"},
            "List Paragraph": {"color": "000000", "sz": "24", "line": "480"},
        }

        for style in styles_element.findall(f".//{{{w_ns}}}style"):
            name_el = style.find(f"{{{w_ns}}}name")
            name_val = name_el.attrib.get(f"{{{w_ns}}}val") if name_el is not None else ""

            cfg = heading_configs.get(name_val)
            if cfg:
                rPr = style.get_or_add_rPr()
                pPr = style.get_or_add_pPr()

                # Fuente
                rFonts = rPr.find(f"{{{w_ns}}}rFonts")
                if rFonts is None:
                    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
                    rPr.append(rFonts)
                # F-10: Eliminar referencias a tema institucional / Calibri / Aptos
                for theme_attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                    theme_qn = qn(f'w:{theme_attr}')
                    if theme_qn in rFonts.attrib:
                        del rFonts.attrib[theme_qn]
                rFonts.set(qn("w:ascii"), font_name)
                rFonts.set(qn("w:hAnsi"), font_name)

                # Color
                color = rPr.find(f"{{{w_ns}}}color")
                if color is None:
                    color = parse_xml(f'<w:color {nsdecls("w")} w:val="{cfg["color"]}"/>')
                    rPr.append(color)
                else:
                    color.set(qn("w:val"), cfg["color"])

                # Tamaño
                sz = rPr.find(f"{{{w_ns}}}sz")
                if sz is None:
                    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{cfg["sz"]}"/>')
                    rPr.append(sz)
                else:
                    sz.set(qn("w:val"), cfg["sz"])

                # Negrita / Cursiva
                if cfg.get("bold"):
                    if rPr.find(f"{{{w_ns}}}b") is None:
                        rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
                if cfg.get("italic"):
                    if rPr.find(f"{{{w_ns}}}i") is None:
                        rPr.append(parse_xml(f'<w:i {nsdecls("w")}/>'))

                # Interlineado y alineacion
                if "line" in cfg:
                    spacing = pPr.find(f"{{{w_ns}}}spacing")
                    if spacing is None:
                        spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:line="{cfg["line"]}" w:lineRule="auto" w:before="0" w:after="0"/>')
                        pPr.append(spacing)
                    else:
                        spacing.set(qn("w:line"), cfg["line"])
                        spacing.set(qn("w:lineRule"), "auto")

                # Tipografía avanzada: Widow Control y Keep With Next (Fase 8)
                widow = pPr.find(f"{{{w_ns}}}widowControl")
                if widow is None:
                    pPr.append(parse_xml(f'<w:widowControl {nsdecls("w")}/>'))

                if name_val.startswith("Heading") or name_val.startswith("Título"):
                    keep_next = pPr.find(f"{{{w_ns}}}keepNext")
                    if keep_next is None:
                        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))

                if "jc" in cfg:
                    jc = pPr.find(f"{{{w_ns}}}jc")
                    if jc is None:
                        jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="{cfg["jc"]}"/>')
                        pPr.append(jc)
                    else:
                        jc.set(qn("w:val"), cfg["jc"])

    except Exception as e:
        print(f"[WARN] Error actualizando styles.xml: {e}")


def _normalize_paragraph_runs(para, target_font: str, target_size) -> None:
    """Normaliza fuente y tamano en todos los runs de un parrafo."""
    for run in para.runs:
        run.font.name = target_font
        run.font.size = target_size


def format_heading_paragraph(p, level: int, text: str, rules: APARuleSet, preserve_text: bool = False) -> None:
    """
    Formatea un titulo estrictamente segun los 5 niveles APA 7:
    - Nivel 1: Centrado, Negrita, Caso Titulo
    - Nivel 2: Izquierda, Negrita, Caso Titulo
    - Nivel 3: Izquierda, Negrita y Cursiva, Caso Titulo
    - Nivel 4: Sangria 1.27cm, Negrita, Termina en punto. Texto en la misma linea
    - Nivel 5: Sangria 1.27cm, Negrita y Cursiva, Termina en punto. Texto en la misma linea
    """
    if not preserve_text:
        p.text = ""  # Limpiar runs
    # F-03: Asignar el estilo nativo de Word ('Heading X' o 'Título X') para enlazar con Navigation Pane y TOC
    try:
        style_candidates = [f"Heading {level}", f"Título {level}"]
        doc_styles = p.part.document.styles
        for sc in style_candidates:
            if sc in doc_styles:
                p.style = doc_styles[sc]
                break
    except Exception:
        try:
            p.style = None
        except Exception:
            pass

    p.paragraph_format.line_spacing = rules.line_spacing
    p.paragraph_format.space_before = Pt(rules.space_before_pt)
    p.paragraph_format.space_after = Pt(rules.space_after_pt)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    if level == 1:
        p.paragraph_format.page_break_before = True

    if preserve_text:
        for r in p.runs:
            set_run_font(r, rules.font_family, rules.font_size_pt)
        return

    if level == 1:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        run.bold = True
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)

    elif level == 2:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        run.bold = True
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)

    elif level == 3:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        run.bold = True
        run.italic = True
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)

    elif level == 4:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(rules.paragraph_indent_cm / 2.54)
        formatted_text: str = text if text.endswith(".") else text + "."
        run = p.add_run(formatted_text + " ")
        run.bold = True
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)

    elif level == 5:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(rules.paragraph_indent_cm / 2.54)
        formatted_text = text if text.endswith(".") else text + "."
        run = p.add_run(formatted_text + " ")
        run.bold = True
        run.italic = True
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)


def format_normal_paragraph(p, text: str, rules: APARuleSet, preserve_text: bool = False) -> None:
    """
    Formatea un parrafo normal APA 7:
    - Fuente uniforme (ej. Times New Roman 12pt)
    - Alineado a la izquierda sin justificar
    - Interlineado 2.0 (doble)
    - Sangria en primera linea de 1.27 cm (0.5 in)
    - Sin espacio entre parrafos (0 pt antes y despues)
    Si preserve_text es True, no destruye el XML interno (usado para proteger ecuaciones y campos).
    """
    if not preserve_text:
        p.text = ""
    # ponytail: clear inherited style to prevent double indentation
    try:
        p.style = None
    except Exception:
        pass
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = rules.line_spacing
    p.paragraph_format.first_line_indent = Inches(rules.paragraph_indent_cm / 2.54)
    p.paragraph_format.space_before = Pt(rules.space_before_pt)
    p.paragraph_format.space_after = Pt(rules.space_after_pt)

    if not preserve_text:
        run = p.add_run(text)
        set_run_font(run, rules.font_family, rules.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)
    else:
        for r in p.runs:
            set_run_font(r, rules.font_family, rules.font_size_pt)


def format_block_quote(p, text: str, rules: APARuleSet) -> None:
    """
    Formatea una cita en bloque (mas de 40 palabras):
    - Todo el bloque sangrado 1.27 cm desde la izquierda
    - Interlineado 2.0
    - Sin comillas alrededor
    """
    p.text = ""
    # ponytail: clear inherited style to prevent double indentation
    try:
        p.style = None
    except Exception:
        pass
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(rules.paragraph_indent_cm / 2.54)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = rules.line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run(text)
    set_run_font(run, rules.font_family, rules.font_size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)


def normalize_global_body_spacing(doc: docx.Document, rules: APARuleSet, cover_paragraph_count: int = 0) -> None:
    """
    Pasada final global: Forzar interlineado doble (w:line=480, lineRule=auto) y 0pt antes/despues
    en el 100% de los parrafos del cuerpo fuera de la portada y celdas de tabla.
    """
    paragraphs = doc.paragraphs[cover_paragraph_count:]
    for p in paragraphs:
        if not p.text.strip():
            continue
        p.paragraph_format.line_spacing = rules.line_spacing
        p.paragraph_format.space_before = Pt(rules.space_before_pt)
        p.paragraph_format.space_after = Pt(rules.space_after_pt)

        # Limpiar negritas residuales en párrafos de cuerpo (no headings)
        style_name = (p.style.name or '') if p.style else ''
        if 'Heading' not in style_name and 'Título' not in style_name:
            all_runs_bold = all(r.bold is True for r in p.runs if r.text.strip())
            if all_runs_bold and p.runs:
                for run in p.runs:
                    if run.bold is True:
                        run.bold = False  # Apagar negrita forzada explícitamente
