"""Traductor de Rúbrica (motor oculto).

Toma una rúbrica (Word o Excel) y un documento (``.docx``) y produce un
informe de cumplimiento: por cada criterio, su peso, un puntaje determinista
cuando es medible y la evidencia que lo respalda.

Honestidad ante todo: los criterios que no se pueden medir sin IA (ortografía,
originalidad, claridad) NO reciben un puntaje inventado; se marcan como
``requiere_ia``. Todo lo determinista cuesta 0 tokens.
"""

from __future__ import annotations

import re
import unicodedata
import uuid
from pathlib import Path
from typing import Any, Optional, Union

# ── utilidades ────────────────────────────────────────────────────────────────


def _norm(text: str) -> str:
    """Minúsculas sin acentos, para comparar nombres de criterio/sección."""
    if not text:
        return ""
    t = unicodedata.normalize("NFKD", text)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return t.lower().strip()


def _to_str(value: Any) -> str:
    return "" if value is None else str(value).strip()


_NUM_RE = re.compile(r"(\d+(?:[.,]\d+)?)")


def _parse_weight(raw: Any) -> Optional[float]:
    """Extrae un peso en % de una celda. Acepta ``'20%'``, ``'20'``, ``'0.2'``.

    Heurística documentada: sin signo ``%``, un valor menor que 1 se interpreta
    como fracción (``0.2`` -> 20).
    """
    text = _to_str(raw)
    if not text:
        return None
    m = _NUM_RE.search(text.replace(",", "."))
    if not m:
        return None
    value = float(m.group(1))
    if "%" in text:
        return value
    if 0 < value < 1:
        return value * 100.0
    return value


def _normalize_weights(criteria: list[dict]) -> None:
    """Normaliza los pesos a un total de 100. Sin pesos -> pesos iguales."""
    total = sum(c["weight"] for c in criteria)
    if total <= 0:
        n = len(criteria) or 1
        for c in criteria:
            c["weight"] = round(100.0 / n, 2)
        return
    if abs(total - 100.0) > 0.5:
        for c in criteria:
            c["weight"] = round(c["weight"] * 100.0 / total, 2)


# ── lectura de la rúbrica ─────────────────────────────────────────────────────

_NAME_HINTS = ("criterio", "criterios", "aspecto", "indicador", "elemento", "nombre", "item")
_WEIGHT_HINTS = ("peso", "puntaje", "valor", "porcentaje", "puntos", "weight", "score", "%")
_DESC_HINTS = ("descripcion", "detalle", "observacion", "descrip")


def _find_header(rows: list[list[str]]) -> Optional[int]:
    for idx, row in enumerate(rows):
        cells = [_norm(_to_str(c)) for c in row]
        has_name = any(any(h in c for h in _NAME_HINTS) for c in cells if c)
        has_weight = any(any(h in c for h in _WEIGHT_HINTS) for c in cells if c)
        if has_name and has_weight:
            return idx
    return None


def _criteria_from_rows(rows: list[list[Any]]) -> list[dict]:
    rows = [[_to_str(c) for c in r] for r in rows]
    header_idx = _find_header(rows)
    if header_idx is not None:
        header = [_norm(c) for c in rows[header_idx]]
        name_col = next((i for i, c in enumerate(header) if any(h in c for h in _NAME_HINTS)), 0)
        weight_col = next((i for i, c in enumerate(header) if any(h in c for h in _WEIGHT_HINTS)), None)
        desc_col = next((i for i, c in enumerate(header) if any(h in c for h in _DESC_HINTS)), None)
        body = rows[header_idx + 1:]
    else:
        name_col, weight_col, desc_col = 0, None, None
        body = rows

    criteria: list[dict] = []
    for row in body:
        name = row[name_col] if name_col < len(row) else ""
        if not name:
            name = next((c for c in row if c), "")
        if not name:
            continue

        weight: Optional[float] = None
        if weight_col is not None and weight_col < len(row):
            weight = _parse_weight(row[weight_col])
        if weight is None:
            for cell in row:
                if "%" in cell:
                    weight = _parse_weight(cell)
                    break
        if weight is None:
            for cell in reversed(row):
                weight = _parse_weight(cell)
                if weight is not None:
                    break

        desc = row[desc_col] if (desc_col is not None and desc_col < len(row)) else ""
        criteria.append({"name": name, "weight": float(weight or 0.0), "description": desc})
    return criteria


_PCT_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")


def _criteria_from_lines(lines: list[str]) -> list[dict]:
    criteria: list[dict] = []
    for line in lines:
        text = _to_str(line)
        if not text:
            continue
        m = _PCT_RE.search(text)
        if not m:
            continue
        name = text[: m.start()].strip(" -–:•\t.")
        if not name:
            continue
        criteria.append(
            {"name": name, "weight": float(m.group(1).replace(",", ".")), "description": ""}
        )
    return criteria


def _rows_and_lines_from_docx(path: Path) -> tuple[list[list[str]], list[str]]:
    import docx  # python-docx

    document = docx.Document(str(path))
    rows: list[list[str]] = []
    for table in document.tables:
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
    if rows:
        return rows, []
    return [], [p.text for p in document.paragraphs]


def parse_rubric(source: Union[str, Path]) -> list[dict]:
    """Lee una rúbrica ``.docx`` o ``.xlsx`` y devuelve sus criterios.

    Cada criterio: ``{"name", "weight", "description"}``. Los pesos se
    normalizan a 100.
    """
    path = Path(source)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from content.ingest import xlsx_to_tables

        rows: list[list[str]] = []
        for table in xlsx_to_tables(path):
            if table.get("headers"):
                rows.append(list(table["headers"]))
            rows.extend([list(r) for r in table.get("rows", [])])
        criteria = _criteria_from_rows(rows)
    elif suffix == ".docx":
        rows, lines = _rows_and_lines_from_docx(path)
        criteria = _criteria_from_rows(rows) if rows else _criteria_from_lines(lines)
    else:
        raise ValueError(
            f"Formato de rúbrica no soportado: {suffix or '(sin extensión)'}. Usa .docx o .xlsx"
        )
    _normalize_weights(criteria)
    return criteria


# ── evaluación determinista del documento ─────────────────────────────────────


def _etype(element: Any) -> str:
    t = getattr(element, "type", "")
    return str(getattr(t, "value", t) or "")


def _has_cover(doc: Any) -> bool:
    portada = getattr(doc, "portada", {}) or {}
    if portada.get("title") or portada.get("detected"):
        return True
    for e in getattr(doc, "elements", []):
        if getattr(e, "is_cover_section", False) or _etype(e) == "portada_block":
            return True
    return False


def _document_stats(doc: Any) -> dict:
    headings: list[str] = []
    parts: list[str] = []
    n_tables = n_images = 0
    for e in getattr(doc, "elements", []):
        etype = _etype(e)
        text = getattr(e, "text", "") or ""
        if text:
            parts.append(text)
        if etype == "heading":
            headings.append(_norm(text))
        elif etype == "table":
            n_tables += 1
        elif etype == "image":
            n_images += 1
    full = "\n".join(parts)
    return {
        "norm_headings": headings,
        "norm_text": _norm(full),
        "words": len(re.findall(r"\w+", full)),
        "n_tables": n_tables,
        "n_images": n_images,
        "n_references": len(getattr(doc, "referencias", []) or []),
        "n_citations": len(getattr(doc, "citas_intext", []) or []),
        "has_cover": _has_cover(doc),
    }


def _result(score: Optional[float], evidence: str, missing: str = "") -> dict:
    if score is None:
        return {"score": None, "status": "requiere_ia", "evidence": evidence, "missing": missing}
    if score >= 85:
        status = "cumple"
    elif score >= 50:
        status = "parcial"
    else:
        status = "no_cumple"
    return {"score": round(score, 1), "status": status, "evidence": evidence, "missing": missing}


def _check_apa_format(doc: Any) -> dict:
    try:
        from modules.proactive_auditor import audit_elements

        findings = audit_elements(list(getattr(doc, "elements", [])))
    except Exception as exc:  # pragma: no cover - defensivo
        return _result(None, f"No se pudo auditar el formato APA: {exc}")
    count = len(findings)
    if count == 0:
        return _result(100, "Sin hallazgos de formato APA")
    kinds = sorted(
        {str(f.get("kind") or f.get("rule") or f.get("type") or "?") for f in findings}
    )
    return _result(
        100 - min(90, count * 8),
        f"{count} hallazgo(s): " + ", ".join(kinds[:6]),
        "Revisa los hallazgos de formato",
    )


def _check_criterion(name: str, stats: dict, doc: Any) -> dict:
    n = _norm(name)

    def has(*keys: str) -> bool:
        return any(k in n for k in keys)

    def heading_has(*keys: str) -> bool:
        return any(any(k in h for k in keys) for h in stats["norm_headings"])

    if has("portada", "caratula", "cover"):
        ok = stats["has_cover"]
        return _result(
            100 if ok else 0,
            "Portada presente" if ok else "No se detectó portada",
            "" if ok else "Agrega la portada",
        )

    if has("introduccion"):
        return _section_result(heading_has("introduccion"), "Introducción")
    if has("objetivo"):
        ok = "objetivo" in stats["norm_text"]
        return _result(
            100 if ok else 0,
            "Objetivos redactados" if ok else "No se encontró la palabra 'objetivo'",
            "" if ok else "Redacta los objetivos",
        )
    if has("marco teorico", "marco", "fundamentacion", "antecedentes"):
        return _section_result(
            heading_has("marco", "fundamentacion", "antecedente"), "Marco teórico"
        )
    if has("metodolog", "materiales y metodos", "metodo"):
        return _section_result(heading_has("metodolog", "materiales", "metodo"), "Metodología")
    if has("resultado"):
        return _section_result(heading_has("resultado"), "Resultados")
    if has("discusion"):
        return _section_result(heading_has("discusion"), "Discusión")
    if has("conclusion"):
        return _section_result(heading_has("conclusion"), "Conclusiones")
    if has("recomendacion"):
        return _section_result(heading_has("recomendacion"), "Recomendaciones")
    if has("resumen", "abstract"):
        ok = stats["words"] > 0 and ("resumen" in stats["norm_text"] or "abstract" in stats["norm_text"])
        return _result(
            100 if ok else 0,
            "Resumen presente" if ok else "No se encontró resumen",
            "" if ok else "Agrega el resumen",
        )

    if has("referencia", "bibliografia", "fuentes"):
        count = stats["n_references"]
        return _result(
            100 if count > 0 else 0,
            f"{count} referencia(s)",
            "" if count else "Agrega las referencias en APA 7",
        )
    if has("cita"):
        count = stats["n_citations"]
        return _result(
            100 if count > 0 else 0,
            f"{count} cita(s) en el texto",
            "" if count else "Agrega citas en el texto",
        )
    if has("tabla", "cuadro"):
        count = stats["n_tables"]
        return _result(
            100 if count > 0 else 0,
            f"{count} tabla(s)",
            "" if count else "Agrega tablas con formato APA",
        )
    if has("figura", "grafico", "imagen"):
        count = stats["n_images"]
        return _result(
            100 if count > 0 else 0,
            f"{count} figura(s)",
            "" if count else "Agrega figuras con formato APA",
        )

    if has("formato", "apa", "presentacion", "normas", "estructura", "estilo"):
        return _check_apa_format(doc)

    if has("ortografia", "redaccion", "gramatica", "claridad", "coherencia"):
        return _result(None, "Requiere revisión de IA (motor probabilístico)")
    if has("originalidad", "plagio", "similitud", "detector"):
        return _result(None, "Requiere detector de IA (motor probabilístico)")

    if has("extension", "pagina", "palabra", "longitud"):
        return _result(None, f"{stats['words']} palabras (sin objetivo explícito)")

    return _result(None, "Criterio no reconocido por los evaluadores deterministas")


def _section_result(present: bool, label: str) -> dict:
    return _result(
        100 if present else 0,
        f"Sección {label}: presente" if present else f"Sección {label}: ausente",
        "" if present else f"Agrega la sección {label}",
    )


def evaluate_rubric(criteria: list[dict], doc: Any) -> tuple[list[dict], dict]:
    """Puntúa cada criterio contra el documento y resume el cumplimiento."""
    stats = _document_stats(doc)
    results: list[dict] = []
    for c in criteria:
        check = _check_criterion(c["name"], stats, doc)
        results.append(
            {
                "name": c["name"],
                "weight": c["weight"],
                "description": c.get("description", ""),
                **check,
            }
        )

    scored = [r for r in results if r["score"] is not None]
    total_weight = sum(r["weight"] for r in scored)
    score = (
        round(sum(r["score"] * r["weight"] for r in scored) / total_weight, 1)
        if total_weight > 0
        else None
    )
    summary = {
        "score": score,
        "coverage_pct": round(sum(r["weight"] for r in scored), 1),
        "scored_criteria": len(scored),
        "total_criteria": len(results),
    }
    return results, summary


def build_rubric_report(
    rubric_path: Union[str, Path],
    docx_path: Union[str, Path],
    *,
    storage_dir: Optional[Union[str, Path]] = None,
    session_id: Optional[str] = None,
) -> dict:
    """Informe de cumplimiento de una rúbrica sobre un documento.

    Determinista y sin IA. Devuelve ``{"rubric", "document", "criteria", "summary"}``.
    """
    from config import STORAGE_DIR
    from parsing.docx_parser import parse_docx_bytes

    storage = Path(storage_dir) if storage_dir else Path(STORAGE_DIR)
    sid = session_id or uuid.uuid4().hex

    criteria = parse_rubric(rubric_path)
    src = Path(docx_path)
    if not src.exists():
        raise FileNotFoundError(f"No existe el .docx: {src}")

    doc = parse_docx_bytes(src.read_bytes(), src.name, sid, storage)
    results, summary = evaluate_rubric(criteria, doc)
    return {
        "rubric": {"source": str(Path(rubric_path)), "criteria_count": len(criteria)},
        "document": {"source": str(src), "session_id": sid},
        "criteria": results,
        "summary": summary,
    }
