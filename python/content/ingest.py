"""Ingesta de fuentes de contenido -> payloads de tablas APA.

Convierte hojas de un .xlsx en dicts con la forma de ``TableSpec``
(``caption`` / ``headers`` / ``rows``) para que la API de contenido los
consume sin conocer ``python-docx`` ni ``openpyxl``. Módulo puro: no
importa ``models`` ni toca el almacén.
"""

from __future__ import annotations

import io
from os import PathLike
from typing import Any, Union


def _require_openpyxl():
    """Import perezoso de openpyxl con error claro si falta."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl no está instalado y es requerido para leer .xlsx. "
            "Instálalo con: pip install openpyxl"
        ) from exc
    return openpyxl


def _cell_str(value: Any) -> str:
    """Todo valor de celda -> str; None -> \"\"."""
    return "" if value is None else str(value)


def _is_blank(text: str) -> bool:
    return text.strip() == ""


def _normalize_sheets(sheets):
    """None -> None (todas); str -> [str]; iterable -> lista de str."""
    if sheets is None:
        return None
    if isinstance(sheets, str):
        return [sheets]
    return [str(s) for s in sheets]


def xlsx_to_tables(
    source: Union[str, PathLike, bytes, bytearray, memoryview],
    *,
    sheets=None,
    header_row: int = 0,
    caption_prefix: str = "Tabla",
) -> list[dict]:
    """Convierte hojas de un .xlsx en tablas APA.

    Args:
        source: Ruta (str/Path) o bytes del archivo .xlsx.
        sheets: ``None`` (todas las hojas), un nombre o una lista de nombres.
            Un nombre inexistente lanza ``KeyError`` con las hojas disponibles.
        header_row: Índice 0-based de la fila de encabezados. Si queda fuera
            de rango o apunta a una fila vacía, ``headers`` queda vacío.
        caption_prefix: Prefijo del caption numerado ("Tabla 1", "Tabla 2", ...).

    Returns:
        Lista de dicts ``{"sheet", "caption", "headers", "rows"}`` con todos
        los valores como ``str`` (``None`` -> ``""``), sin filas totalmente
        vacías, filas rectangulares y columnas finales vacías recortadas.
        El caption se numera de forma secuencial (1-based) sobre las tablas
        devueltas.
    """
    openpyxl = _require_openpyxl()

    if isinstance(source, (bytes, bytearray, memoryview)):
        workbook = openpyxl.load_workbook(
            io.BytesIO(bytes(source)), data_only=True, read_only=True
        )
    else:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)

    try:
        wanted = _normalize_sheets(sheets)
        if wanted is None:
            wanted = list(workbook.sheetnames)

        tables: list[dict] = []
        for sheet_name in wanted:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(
                    f"La hoja {sheet_name!r} no existe en el workbook. "
                    f"Hojas disponibles: {workbook.sheetnames}"
                )
            worksheet = workbook[sheet_name]
            raw_rows = [
                [_cell_str(v) for v in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            # Filas con contenido (sin importar su posición).
            content_rows = [r for r in raw_rows if not all(_is_blank(c) for c in r)]
            if not content_rows:
                continue  # hoja sin contenido -> no genera tabla

            if (
                0 <= header_row < len(raw_rows)
                and not all(_is_blank(c) for c in raw_rows[header_row])
            ):
                headers = list(raw_rows[header_row])
            else:
                headers = []
            body = [
                r
                for i, r in enumerate(raw_rows)
                if i > header_row and not all(_is_blank(c) for c in r)
            ]
            if not headers and not body:
                continue

            # Rectangularizar filas y recortar columnas finales vacías.
            width = max([len(headers)] + [len(r) for r in body])
            headers = headers + [""] * (width - len(headers))
            body = [r + [""] * (width - len(r)) for r in body]
            while width > 0 and _is_blank(headers[width - 1]) and all(
                _is_blank(r[width - 1]) for r in body
            ):
                width -= 1

            tables.append(
                {
                    "sheet": sheet_name,
                    "caption": f"{caption_prefix} {len(tables) + 1}",
                    "headers": headers[:width],
                    "rows": [r[:width] for r in body],
                }
            )
        return tables
    finally:
        workbook.close()
